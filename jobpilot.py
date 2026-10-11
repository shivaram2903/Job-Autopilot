"""Core services for JobPilot AI, a local-first job search and application tracker."""
from __future__ import annotations

import html
import json
import os
import re
import smtplib
import sqlite3
from datetime import datetime, timedelta, timezone
from email.message import EmailMessage
from pathlib import Path
from typing import Any, Iterable
from zoneinfo import ZoneInfo

import requests

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
RESUME_DIR = DATA_DIR / "resumes"
DB_PATH = Path(os.getenv("JOBPILOT_DB_PATH", str(DATA_DIR / "jobpilot.sqlite3")))
LOCAL_TZ = ZoneInfo(os.getenv("JOBPILOT_TIMEZONE", "Asia/Kolkata"))

REMOTIVE_URL = "https://remotive.com/api/remote-jobs"
ARBEITNOW_URL = "https://www.arbeitnow.com/api/job-board-api"
REQUEST_TIMEOUT = 20

APPLICATION_STATUSES = [
    "Discovered", "Shortlisted", "Awaiting Approval", "Application in Progress",
    "Submitted", "Application Failed", "Assessment Pending", "Interview Scheduled",
    "Under Review", "Rejected", "Offer Received", "Withdrawn", "Status Unknown",
]

KNOWN_SKILLS = [
    "python", "sql", "mysql", "postgresql", "azure", "aws", "gcp", "linux",
    "git", "github", "docker", "kubernetes", "jenkins", "ansible", "terraform",
    "pandas", "numpy", "scikit-learn", "tensorflow", "pytorch", "machine learning",
    "deep learning", "artificial intelligence", "data analysis", "power bi", "tableau",
    "excel", "java", "javascript", "typescript", "react", "node.js", "fastapi",
    "django", "flask", "rest api", "etl", "data engineering", "devops", "azure monitor",
    "query optimization", "performance tuning", "database administration", "postgresql flexible server",
    "customer support", "technical support", "troubleshooting", "incident management",
]


def utc_now() -> datetime:
    return datetime.now(timezone.utc)


def iso_now() -> str:
    return utc_now().isoformat(timespec="seconds")


def get_conn() -> sqlite3.Connection:
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH, timeout=30)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("PRAGMA journal_mode = WAL")
    return conn


def init_db() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    RESUME_DIR.mkdir(parents=True, exist_ok=True)
    with get_conn() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS profile (
                id INTEGER PRIMARY KEY CHECK (id = 1),
                full_name TEXT NOT NULL DEFAULT '',
                email TEXT NOT NULL DEFAULT '',
                phone TEXT NOT NULL DEFAULT '',
                skills TEXT NOT NULL DEFAULT '',
                experience_years REAL NOT NULL DEFAULT 0,
                education TEXT NOT NULL DEFAULT '',
                target_roles TEXT NOT NULL DEFAULT '',
                preferred_location TEXT NOT NULL DEFAULT '',
                remote_only INTEGER NOT NULL DEFAULT 0,
                employment_types TEXT NOT NULL DEFAULT 'Full-time',
                resume_text TEXT NOT NULL DEFAULT '',
                resume_path TEXT NOT NULL DEFAULT '',
                updated_at TEXT NOT NULL DEFAULT ''
            );
            INSERT OR IGNORE INTO profile (id) VALUES (1);

            CREATE TABLE IF NOT EXISTS jobs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                external_id TEXT NOT NULL,
                source TEXT NOT NULL,
                title TEXT NOT NULL,
                company TEXT NOT NULL DEFAULT '',
                location TEXT NOT NULL DEFAULT '',
                description TEXT NOT NULL DEFAULT '',
                url TEXT NOT NULL,
                salary TEXT NOT NULL DEFAULT '',
                job_type TEXT NOT NULL DEFAULT '',
                published_at TEXT NOT NULL DEFAULT '',
                source_attribution TEXT NOT NULL DEFAULT '',
                fetched_at TEXT NOT NULL,
                UNIQUE(source, external_id),
                UNIQUE(source, url)
            );

            CREATE TABLE IF NOT EXISTS applications (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                job_id INTEGER NOT NULL UNIQUE REFERENCES jobs(id) ON DELETE CASCADE,
                status TEXT NOT NULL DEFAULT 'Shortlisted',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                notes TEXT NOT NULL DEFAULT '',
                next_action TEXT NOT NULL DEFAULT '',
                follow_up_date TEXT NOT NULL DEFAULT '',
                confirmation_evidence TEXT NOT NULL DEFAULT '',
                resume_version TEXT NOT NULL DEFAULT ''
            );

            CREATE TABLE IF NOT EXISTS status_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                application_id INTEGER NOT NULL REFERENCES applications(id) ON DELETE CASCADE,
                previous_status TEXT NOT NULL,
                new_status TEXT NOT NULL,
                changed_at TEXT NOT NULL,
                source TEXT NOT NULL DEFAULT 'manual',
                evidence TEXT NOT NULL DEFAULT ''
            );

            CREATE TABLE IF NOT EXISTS settings (
                key TEXT PRIMARY KEY,
                value TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS reports (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                report_date TEXT NOT NULL,
                content TEXT NOT NULL,
                created_at TEXT NOT NULL,
                sent_at TEXT NOT NULL DEFAULT ''
            );
            """
        )


def row_to_dict(row: sqlite3.Row | dict | None) -> dict[str, Any] | None:
    return dict(row) if row is not None else None


def get_profile() -> dict[str, Any]:
    with get_conn() as conn:
        row = conn.execute("SELECT * FROM profile WHERE id = 1").fetchone()
    profile = dict(row) if row else {}
    profile["remote_only"] = bool(profile.get("remote_only", 0))
    return profile


def save_profile(profile: dict[str, Any]) -> None:
    fields = [
        "full_name", "email", "phone", "skills", "experience_years", "education",
        "target_roles", "preferred_location", "remote_only", "employment_types",
        "resume_text", "resume_path",
    ]
    values = []
    for field in fields:
        value = profile.get(field, "")
        if field == "experience_years":
            try:
                value = max(0.0, float(value or 0))
            except (TypeError, ValueError):
                value = 0.0
        elif field == "remote_only":
            value = int(bool(value))
        else:
            value = str(value or "")
        values.append(value)
    values.extend([iso_now(), 1])
    assignments = ", ".join(f"{field} = ?" for field in fields)
    with get_conn() as conn:
        conn.execute(
            f"UPDATE profile SET {assignments}, updated_at = ? WHERE id = ?",
            values,
        )


def get_setting(key: str, default: str = "") -> str:
    with get_conn() as conn:
        row = conn.execute("SELECT value FROM settings WHERE key = ?", (key,)).fetchone()
    return str(row["value"]) if row else default


def set_setting(key: str, value: str) -> None:
    with get_conn() as conn:
        conn.execute(
            "INSERT INTO settings(key, value) VALUES(?, ?) "
            "ON CONFLICT(key) DO UPDATE SET value = excluded.value",
            (key, str(value)),
        )


def clean_html(value: str) -> str:
    value = html.unescape(value or "")
    value = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", value)
    value = re.sub(r"(?s)<[^>]+>", " ", value)
    value = re.sub(r"\s+", " ", value)
    return value.strip()


def split_values(value: str | Iterable[str] | None) -> list[str]:
    if value is None:
        return []
    if not isinstance(value, str):
        return [str(x).strip() for x in value if str(x).strip()]
    return [part.strip() for part in re.split(r"[,;\n]", value) if part.strip()]


def normalize_text(value: str) -> str:
    return re.sub(r"[^a-z0-9+#.]+", " ", (value or "").lower()).strip()


def extract_skill_keywords(resume_text: str) -> list[str]:
    text = normalize_text(resume_text)
    return [skill for skill in KNOWN_SKILLS if normalize_text(skill) in text]


def parse_resume_bytes(filename: str, content: bytes) -> str:
    """Extract text from PDF/DOCX without uploading the document to a third party."""
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        import io
        reader = PdfReader(io.BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    if ext == ".docx":
        from docx import Document
        import io
        doc = Document(io.BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(paragraphs).strip()
    if ext == ".txt":
        return content.decode("utf-8", errors="replace").strip()
    raise ValueError("Unsupported file type. Please upload a PDF, DOCX, or TXT resume.")


def save_resume_file(filename: str, content: bytes) -> str:
    safe_name = re.sub(r"[^a-zA-Z0-9._-]+", "_", Path(filename).name)[:100] or "resume"
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = RESUME_DIR / f"{stamp}_{safe_name}"
    path.write_bytes(content)
    return str(path)


def _normalize_job(raw: dict[str, Any], source: str) -> dict[str, str]:
    if source == "Remotive":
        external_id = str(raw.get("id") or raw.get("url") or "")
        location = str(raw.get("candidate_required_location") or "Remote")
        published = str(raw.get("publication_date") or "")
        description = clean_html(str(raw.get("description") or ""))
        source_attribution = "Remotive"
    elif source == "Arbeitnow":
        external_id = str(raw.get("slug") or raw.get("url") or "")
        location = str(raw.get("location") or ("Remote" if raw.get("remote") else "Not specified"))
        published = ""
        if raw.get("created_at"):
            try:
                published = datetime.fromtimestamp(int(raw["created_at"]), timezone.utc).isoformat()
            except (TypeError, ValueError, OSError):
                published = ""
        tags = raw.get("tags") or []
        if isinstance(tags, list):
            description = clean_html(str(raw.get("description") or "")) + " Skills/tags: " + ", ".join(map(str, tags))
        else:
            description = clean_html(str(raw.get("description") or ""))
        source_attribution = "Arbeitnow"
    else:
        raise ValueError(f"Unknown job source: {source}")

    title = str(raw.get("title") or "Untitled role").strip()
    company = str(raw.get("company_name") or raw.get("company") or "Company not listed").strip()
    url = str(raw.get("url") or "").strip()
    tags = raw.get("tags") or []
    job_type = str(raw.get("job_type") or ", ".join(tags) if isinstance(tags, list) else raw.get("job_type") or "")
    salary = str(raw.get("salary") or "")
    return {
        "external_id": external_id or url or f"{source}:{title}:{company}",
        "source": source,
        "title": title,
        "company": company,
        "location": location,
        "description": description,
        "url": url,
        "salary": salary,
        "job_type": job_type,
        "published_at": published,
        "source_attribution": source_attribution,
    }


def _upsert_jobs(job_items: list[dict[str, str]]) -> int:
    inserted_or_updated = 0
    now = iso_now()
    with get_conn() as conn:
        for job in job_items:
            if not job.get("url"):
                continue
            conn.execute(
                """INSERT INTO jobs(
                    external_id, source, title, company, location, description, url, salary,
                    job_type, published_at, source_attribution, fetched_at
                ) VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(source, external_id) DO UPDATE SET
                    title=excluded.title, company=excluded.company, location=excluded.location,
                    description=excluded.description, url=excluded.url, salary=excluded.salary,
                    job_type=excluded.job_type, published_at=excluded.published_at,
                    source_attribution=excluded.source_attribution, fetched_at=excluded.fetched_at""",
                (
                    job["external_id"], job["source"], job["title"], job["company"],
                    job["location"], job["description"], job["url"], job["salary"],
                    job["job_type"], job["published_at"], job["source_attribution"], now,
                ),
            )
            inserted_or_updated += 1
    return inserted_or_updated


def _fetch_json(url: str, params: dict[str, Any] | None = None) -> dict[str, Any]:
    response = requests.get(
        url,
        params=params,
        timeout=REQUEST_TIMEOUT,
        headers={"User-Agent": "JobPilotAI/0.1 (personal job-search assistant)"},
    )
    response.raise_for_status()
    payload = response.json()
    if not isinstance(payload, dict):
        raise ValueError("Unexpected response from job source.")
    return payload


def refresh_sources(sources: list[str] | None = None, force: bool = False) -> dict[str, Any]:
    """Fetch public job feeds with local cooldowns to avoid excessive API requests."""
    sources = sources or ["Remotive", "Arbeitnow"]
    outcomes: dict[str, Any] = {}
    profile = get_profile()
    query = " ".join(split_values(profile.get("target_roles", ""))[:2])
    now = utc_now()
    cooldown_hours = {"Remotive": 6, "Arbeitnow": 3}

    for source in sources:
        key = f"last_fetch_{source.lower()}"
        last_value = get_setting(key, "")
        eligible = force or not last_value
        if last_value and not force:
            try:
                last_dt = datetime.fromisoformat(last_value)
                if last_dt.tzinfo is None:
                    last_dt = last_dt.replace(tzinfo=timezone.utc)
                eligible = now - last_dt >= timedelta(hours=cooldown_hours.get(source, 4))
            except ValueError:
                eligible = True
        if not eligible:
            outcomes[source] = {
                "ok": True,
                "cached": True,
                "message": f"Using saved {source} jobs; next refresh is limited to respect feed request guidance.",
                "count": count_jobs(source=source),
            }
            continue

        try:
            if source == "Remotive":
                payload = _fetch_json(REMOTIVE_URL, {"limit": 300})
                rows = payload.get("jobs") or []
                normalized = [_normalize_job(item, source) for item in rows if isinstance(item, dict)]
            elif source == "Arbeitnow":
                payload = _fetch_json(ARBEITNOW_URL)
                rows = payload.get("data") or []
                normalized = [_normalize_job(item, source) for item in rows if isinstance(item, dict)]
            else:
                outcomes[source] = {"ok": False, "cached": False, "message": "Source is not configured."}
                continue
            count = _upsert_jobs(normalized)
            set_setting(key, now.isoformat(timespec="seconds"))
            outcomes[source] = {"ok": True, "cached": False, "message": f"Fetched/updated {count} listings.", "count": count}
        except (requests.RequestException, ValueError, KeyError) as exc:
            outcomes[source] = {
                "ok": False,
                "cached": False,
                "message": f"Could not refresh this source: {exc}",
                "count": count_jobs(source=source),
            }
    return outcomes


def count_jobs(source: str | None = None) -> int:
    with get_conn() as conn:
        if source:
            row = conn.execute("SELECT COUNT(*) AS n FROM jobs WHERE source = ?", (source,)).fetchone()
        else:
            row = conn.execute("SELECT COUNT(*) AS n FROM jobs").fetchone()
    return int(row["n"])


def score_job(profile: dict[str, Any], job: dict[str, Any]) -> dict[str, Any]:
    """Transparent heuristic score. This is not an LLM evaluation."""
    body = normalize_text(" ".join([
        str(job.get("title", "")), str(job.get("description", "")),
        str(job.get("job_type", "")), str(job.get("location", "")),
    ]))
    title = normalize_text(str(job.get("title", "")))
    profile_skills = split_values(profile.get("skills", ""))
    if not profile_skills:
        profile_skills = extract_skill_keywords(str(profile.get("resume_text", "")))
    skills_normalized = [normalize_text(item) for item in profile_skills if normalize_text(item)]
    matched = [original for original, normalized in zip(profile_skills, skills_normalized) if normalized in body]
    missing = [original for original, normalized in zip(profile_skills, skills_normalized) if normalized not in body]
    if skills_normalized:
        skill_points = 55 * len(matched) / len(skills_normalized)
    else:
        skill_points = 27.5

    target_roles = split_values(profile.get("target_roles", ""))
    role_points = 12.5
    role_reason = "No target roles configured"
    if target_roles:
        best = 0.0
        for role in target_roles:
            role_norm = normalize_text(role)
            if role_norm and (role_norm in title or title in role_norm):
                best = max(best, 1.0)
            else:
                role_tokens = set(role_norm.split())
                title_tokens = set(title.split())
                if role_tokens:
                    best = max(best, len(role_tokens & title_tokens) / len(role_tokens))
        role_points = 25 * best
        role_reason = "Target role aligns with job title" if best >= 0.6 else "Partial or weak title alignment"

    preferred_location = normalize_text(str(profile.get("preferred_location", "")))
    job_location = normalize_text(str(job.get("location", "")))
    is_remote = bool(re.search(r"\bremote\b|worldwide|anywhere", body))
    if profile.get("remote_only"):
        location_points = 10 if is_remote else 0
        location_reason = "Remote-compatible" if is_remote else "Remote-only preference not clearly met"
    elif not preferred_location:
        location_points = 7.5
        location_reason = "No location preference configured"
    elif preferred_location in job_location or preferred_location in body:
        location_points = 10
        location_reason = "Preferred location appears in listing"
    elif is_remote:
        location_points = 8
        location_reason = "Remote role may fit location preference"
    else:
        location_points = 2
        location_reason = "Location does not clearly match"

    candidate_years = float(profile.get("experience_years") or 0)
    year_matches = re.findall(r"(\d+)\s*\+?\s*years?", body)
    required_years = min((int(x) for x in year_matches), default=None)
    if required_years is None:
        experience_points = 7.5
        experience_reason = "No clear years-of-experience requirement detected"
    elif candidate_years >= required_years:
        experience_points = 10
        experience_reason = f"Candidate experience meets detected minimum ({required_years} years)"
    elif candidate_years >= max(0, required_years - 1):
        experience_points = 5
        experience_reason = f"Candidate is close to detected minimum ({required_years} years)"
    else:
        experience_points = 0
        experience_reason = f"Listing may require at least {required_years} years"

    score = round(max(0, min(100, skill_points + role_points + location_points + experience_points)))
    if score >= 90:
        category = "Excellent Match"
    elif score >= 75:
        category = "Strong Match"
    elif score >= 60:
        category = "Possible Match"
    else:
        category = "Low Match"
    reasons = [role_reason, location_reason, experience_reason]
    if matched:
        reasons.insert(0, f"Matching skills: {', '.join(matched[:8])}")
    elif skills_normalized:
        reasons.insert(0, "No profile skills were detected in the job text")
    return {
        "score": score,
        "category": category,
        "matched_skills": matched,
        "missing_skills": missing,
        "reasons": reasons,
        "breakdown": {
            "skills": round(skill_points, 1),
            "role": round(role_points, 1),
            "location": round(location_points, 1),
            "experience": round(experience_points, 1),
        },
    }


def get_jobs(search: str = "", min_score: int = 0, sources: list[str] | None = None) -> list[dict[str, Any]]:
    with get_conn() as conn:
        if sources:
            placeholders = ",".join("?" for _ in sources)
            rows = conn.execute(f"SELECT * FROM jobs WHERE source IN ({placeholders}) ORDER BY published_at DESC, fetched_at DESC", sources).fetchall()
        else:
            rows = conn.execute("SELECT * FROM jobs ORDER BY published_at DESC, fetched_at DESC").fetchall()
    profile = get_profile()
    search_groups = [normalize_text(group) for group in split_values(search)]
    search_groups = [group for group in search_groups if group]
    results = []
    for row in rows:
        job = dict(row)
        haystack = normalize_text(" ".join([job["title"], job["company"], job["location"], job["description"], job["source"]]))
        # Comma/newline separated search phrases are alternatives, while words
        # inside one phrase must all be present.
        if search_groups and not any(all(term in haystack for term in group.split()) for group in search_groups):
            continue
        score = score_job(profile, job)
        job.update(score)
        if score["score"] >= min_score:
            results.append(job)
    results.sort(key=lambda item: (item["score"], item.get("published_at", "")), reverse=True)
    return results


def add_application(job_id: int, status: str = "Shortlisted") -> bool:
    if status not in APPLICATION_STATUSES:
        raise ValueError("Invalid application status")
    now = iso_now()
    with get_conn() as conn:
        existing = conn.execute("SELECT id FROM applications WHERE job_id = ?", (job_id,)).fetchone()
        if existing:
            return False
        conn.execute(
            "INSERT INTO applications(job_id, status, created_at, updated_at) VALUES (?, ?, ?, ?)",
            (job_id, status, now, now),
        )
        app_id = conn.execute("SELECT last_insert_rowid() AS id").fetchone()["id"]
        conn.execute(
            "INSERT INTO status_history(application_id, previous_status, new_status, changed_at, source) VALUES (?, ?, ?, ?, ?)",
            (app_id, "", status, now, "created"),
        )
    return True


def update_application(
    application_id: int,
    status: str,
    notes: str = "",
    next_action: str = "",
    follow_up_date: str = "",
    evidence: str = "",
    source: str = "manual",
) -> None:
    if status not in APPLICATION_STATUSES:
        raise ValueError("Invalid application status")
    now = iso_now()
    with get_conn() as conn:
        row = conn.execute("SELECT status FROM applications WHERE id = ?", (application_id,)).fetchone()
        if not row:
            raise ValueError("Application not found")
        previous = str(row["status"])
        conn.execute(
            """UPDATE applications SET status=?, updated_at=?, notes=?, next_action=?,
               follow_up_date=?, confirmation_evidence=? WHERE id=?""",
            (status, now, notes.strip(), next_action.strip(), follow_up_date.strip(), evidence.strip(), application_id),
        )
        if status != previous:
            conn.execute(
                """INSERT INTO status_history(application_id, previous_status, new_status, changed_at, source, evidence)
                   VALUES (?, ?, ?, ?, ?, ?)""",
                (application_id, previous, status, now, source, evidence.strip()),
            )


def get_applications(status_filter: str | None = None, search: str = "") -> list[dict[str, Any]]:
    sql = """SELECT a.*, j.title, j.company, j.location, j.url, j.source, j.salary, j.published_at,
              j.source_attribution, j.id AS job_id
              FROM applications a JOIN jobs j ON j.id = a.job_id"""
    params: list[Any] = []
    where = []
    if status_filter and status_filter != "All":
        where.append("a.status = ?")
        params.append(status_filter)
    if search.strip():
        where.append("(LOWER(j.title) LIKE ? OR LOWER(j.company) LIKE ? OR LOWER(j.location) LIKE ?)")
        like = f"%{search.strip().lower()}%"
        params.extend([like, like, like])
    if where:
        sql += " WHERE " + " AND ".join(where)
    sql += " ORDER BY a.updated_at DESC"
    with get_conn() as conn:
        rows = conn.execute(sql, params).fetchall()
    results = []
    profile = get_profile()
    for row in rows:
        item = dict(row)
        with get_conn() as conn:
            job_row = conn.execute("SELECT * FROM jobs WHERE id = ?", (item["job_id"],)).fetchone()
        if job_row:
            item["match_score"] = score_job(profile, dict(job_row))["score"]
        else:
            item["match_score"] = 0
        results.append(item)
    return results


def get_status_history(application_id: int | None = None, since_iso: str | None = None) -> list[dict[str, Any]]:
    sql = """SELECT h.*, j.title, j.company, j.url FROM status_history h
             JOIN applications a ON a.id = h.application_id JOIN jobs j ON j.id = a.job_id"""
    clauses, params = [], []
    if application_id is not None:
        clauses.append("h.application_id = ?")
        params.append(application_id)
    if since_iso:
        clauses.append("h.changed_at >= ?")
        params.append(since_iso)
    if clauses:
        sql += " WHERE " + " AND ".join(clauses)
    sql += " ORDER BY h.changed_at DESC"
    with get_conn() as conn:
        return [dict(row) for row in conn.execute(sql, params).fetchall()]


def _local_day_start_utc(now: datetime | None = None) -> datetime:
    current_local = (now or datetime.now(LOCAL_TZ)).astimezone(LOCAL_TZ)
    start_local = current_local.replace(hour=0, minute=0, second=0, microsecond=0)
    return start_local.astimezone(timezone.utc)


def create_daily_report() -> dict[str, Any]:
    start_utc = _local_day_start_utc()
    start_iso = start_utc.isoformat(timespec="seconds")
    today_local = datetime.now(LOCAL_TZ).date().isoformat()
    apps = get_applications()
    history = get_status_history(since_iso=start_iso)
    # Count actual status transitions, not merely records edited today.
    submitted_today = [h for h in history if h["new_status"] == "Submitted"]
    awaiting = [a for a in apps if a["status"] == "Awaiting Approval"]
    failures = [h for h in history if h["new_status"] == "Application Failed"]
    interviews = [a for a in apps if a["status"] == "Interview Scheduled"]
    assessments = [a for a in apps if a["status"] == "Assessment Pending"]
    match_threshold = int(get_setting("min_match_score", "60"))
    matching_jobs_count = len(get_jobs(min_score=match_threshold))

    with get_conn() as conn:
        new_jobs_row = conn.execute("SELECT COUNT(*) AS n FROM jobs WHERE fetched_at >= ?", (start_iso,)).fetchone()
        total_tracked = conn.execute("SELECT COUNT(*) AS n FROM applications").fetchone()["n"]
    lines = [
        f"JobPilot AI daily report — {today_local} ({LOCAL_TZ.key})",
        "=" * 58,
        f"Job listings refreshed/added today: {new_jobs_row['n']}",
        f"Saved listings meeting the current match threshold ({match_threshold}+): {matching_jobs_count}",
        f"Applications currently tracked: {total_tracked}",
        f"Applications marked submitted today: {len(submitted_today)}",
        f"Awaiting approval: {len(awaiting)}",
        f"Application failures today: {len(failures)}",
        f"Interviews scheduled (all tracked): {len(interviews)}",
        f"Assessments pending (all tracked): {len(assessments)}",
        f"Status changes today: {len(history)}",
        "",
        "APPLICATIONS SUBMITTED TODAY",
        "-" * 30,
    ]
    if submitted_today:
        for event in submitted_today:
            lines.append(f"• {event['company']} — {event['title']} | Confirmed transition: {event['changed_at']} | {event['url']}")
    else:
        lines.append("No applications were marked submitted today.")
    lines.extend(["", "STATUS CHANGES TODAY", "-" * 30])
    if history:
        for item in history:
            old = item["previous_status"] or "New"
            evidence = f" | Evidence: {item['evidence']}" if item.get("evidence") else ""
            lines.append(f"• {item['company']} — {item['title']}: {old} → {item['new_status']} at {item['changed_at']} ({item['source']}){evidence}")
    else:
        lines.append("No recorded status changes today.")
    lines.extend(["", "PENDING ACTIONS", "-" * 30])
    pending = [a for a in apps if a["status"] in {"Awaiting Approval", "Application Failed", "Assessment Pending", "Interview Scheduled"}]
    if pending:
        for app in pending[:30]:
            action = f" | Next: {app['next_action']}" if app.get("next_action") else ""
            due = f" | Follow-up: {app['follow_up_date']}" if app.get("follow_up_date") else ""
            lines.append(f"• {app['company']} — {app['title']} [{app['status']}]{action}{due}")
    else:
        lines.append("No pending actions recorded.")
    content = "\n".join(lines)
    with get_conn() as conn:
        conn.execute(
            "INSERT INTO reports(report_date, content, created_at) VALUES(?, ?, ?)",
            (today_local, content, iso_now()),
        )
    return {
        "date": today_local,
        "content": content,
        "submitted_today": len(submitted_today),
        "awaiting_approval": len(awaiting),
        "failures_today": len(failures),
        "status_changes_today": len(history),
    }


def send_email(subject: str, body: str, recipient: str | None = None) -> None:
    """Send email using SMTP settings in environment variables."""
    host = os.getenv("SMTP_HOST", "").strip()
    user = os.getenv("SMTP_USER", "").strip()
    password = os.getenv("SMTP_PASSWORD", "")
    sender = os.getenv("SMTP_FROM", user).strip()
    to = (recipient or os.getenv("REPORT_EMAIL_TO", "")).strip()
    port = int(os.getenv("SMTP_PORT", "587"))
    use_tls = os.getenv("SMTP_USE_TLS", "true").lower() not in {"0", "false", "no"}
    if not (host and sender and to):
        raise RuntimeError("Email is not configured. Set SMTP_HOST, SMTP_FROM, and REPORT_EMAIL_TO in .env.")
    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = sender
    msg["To"] = to
    msg.set_content(body)
    with smtplib.SMTP(host, port, timeout=25) as server:
        if use_tls:
            server.starttls()
        if user and password:
            server.login(user, password)
        server.send_message(msg)


def mark_latest_report_sent() -> None:
    with get_conn() as conn:
        row = conn.execute("SELECT id FROM reports ORDER BY id DESC LIMIT 1").fetchone()
        if row:
            conn.execute("UPDATE reports SET sent_at = ? WHERE id = ?", (iso_now(), row["id"]))


def get_recent_reports(limit: int = 10) -> list[dict[str, Any]]:
    with get_conn() as conn:
        rows = conn.execute("SELECT * FROM reports ORDER BY id DESC LIMIT ?", (limit,)).fetchall()
    return [dict(row) for row in rows]
